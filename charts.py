"""
시드 키워드별 tt/gg/az 오버레이 차트 엑셀 생성 (틱톡 데이터 구간만).

각 시트 = 키워드 1개: 주간 값 + 틱톡 스파이크 onset 마커(spike_events 결과).

사용: python charts.py [LANG]   (기본 US)
출력: data/seed_charts_{LANG}.xlsx
"""
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference

HERE = Path(__file__).parent


def add_sheet(wb: Workbook, kw: str, w: pd.DataFrame, onset_dates: set[str]) -> None:
    ws = wb.create_sheet(kw[:31])
    ws.append(["date", "tt", "gg", "az", "tt_onset"])
    for _, r in w.iterrows():
        d = r["date"].date().isoformat()
        ws.append([d, r["tt"], r["gg"], r["az"], r["tt"] if d in onset_dates else None])
    n = len(w) + 1

    ch = LineChart()
    ch.title = f"{kw} — TikTok / Google / Amazon (주간, 0-100)"
    ch.height, ch.width = 9, 28
    ch.y_axis.title, ch.x_axis.title = "정규화 값", "주(토요일 종료)"
    ch.add_data(Reference(ws, min_col=2, max_col=5, min_row=1, max_row=n), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n))
    for s in ch.series[:3]:
        s.smooth = False
    mk = ch.series[3]                       # tt_onset: 선 없이 마커만
    mk.graphicalProperties.line.noFill = True
    mk.marker.symbol, mk.marker.size = "triangle", 9
    ws.add_chart(ch, "G2")


def main() -> None:
    lang = sys.argv[1].upper() if len(sys.argv) > 1 else "US"
    df = pd.read_csv(HERE / "data" / f"series_{lang}.csv", parse_dates=["date"])
    df = df.sort_values(["keyword", "date"]).reset_index(drop=True)
    ev = pd.read_csv(HERE / "data" / f"spike_events_{lang}.csv")
    onset_by_kw = ev.groupby("keyword")["tt_onset"].apply(set).to_dict()

    wb = Workbook()
    wb.remove(wb.active)
    for kw, g in df.groupby("keyword", sort=True):
        g = g.reset_index(drop=True)
        tt_idx = g.index[g["tt"] > 0]
        if len(tt_idx) == 0:
            continue
        w = g.loc[tt_idx.min(): tt_idx.max()]
        add_sheet(wb, kw, w, onset_by_kw.get(kw, set()))
    path = HERE / "data" / f"seed_charts_{lang}.xlsx"
    wb.save(path)
    print(f"저장: {path} (시트 {len(wb.sheetnames)}개)")


if __name__ == "__main__":
    main()
